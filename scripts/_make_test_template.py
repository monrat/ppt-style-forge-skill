#!/usr/bin/env python3
"""Build a small BRAND-FLAVOURED test template for end-to-end self-tests.

This is NOT a shipped skill asset — it exists so the generate->validate->render
pipeline can be exercised without a real brand .pptx. It produces a template
with a custom theme (so off-brand colour checks are meaningful), a master
footer, and the standard set of named layouts that python-pptx exposes.

    python3 scripts/_make_test_template.py tests/fixtures/sample.pptx

The leading underscore marks it as an internal/test-only script.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


# A small but distinct palette. None of these match the Office-default accent
# set, so a deck that introduces Office defaults will be flagged off-brand.
BRAND_PRIMARY = RGBColor(0x0B, 0x3D, 0x91)     # deep blue
BRAND_ACCENT = RGBColor(0xE8, 0x71, 0x22)      # warm orange
BRAND_DARK = RGBColor(0x1A, 0x1A, 0x2E)        # near-black navy
BRAND_LIGHT = RGBColor(0xF5, 0xF5, 0xF7)       # off-white


def _rewrite_theme_xml(prs: Presentation) -> None:
    """Patch theme1.xml: colour scheme + fonts."""
    import re
    from pptx.oxml.ns import qn as _qn

    # Access the theme part through the first slide master's relationships.
    master = prs.slide_masters[0]
    theme_part = None
    for rel in master.part.rels.values():
        if "theme" in rel.reltype:
            theme_part = rel.target_part
            break
    if theme_part is None:
        return
    blob = theme_part.blob.decode("utf-8")

    # Replace the colour scheme's srgbClr values with brand colours. Map by the
    # surrounding element name (dk1/lt1/dk2/lt2/accent1..6/hlink/folHlink).
    palette = {
        "dk1": BRAND_DARK, "lt1": BRAND_LIGHT,
        "dk2": BRAND_DARK, "lt2": BRAND_LIGHT,
        "accent1": BRAND_PRIMARY, "accent2": BRAND_ACCENT,
        "accent3": BRAND_PRIMARY, "accent4": BRAND_ACCENT,
        "accent5": BRAND_PRIMARY, "accent6": BRAND_ACCENT,
        "hlink": BRAND_PRIMARY, "folHlink": BRAND_PRIMARY,
    }
    for name, rgb in palette.items():
        # match e.g. <a:dk1>...<a:srgbClr val="xxxxxx"/>...</a:dk1>
        pat = re.compile(
            r"(<a:" + name + r">.*?<a:srgbClr val=\")([0-9A-Fa-f]{6})(\".*?</a:" + name + r">)",
            re.DOTALL,
        )
        hexval = "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])
        blob = pat.sub(lambda m: m.group(1) + hexval + m.group(3), blob)

    # Set major/minor fonts to a non-default pairing so theme-font checks differ
    # from the stock "Calibri".
    blob = re.sub(
        r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")',
        r'\1Mercury Display\2', blob)
    blob = re.sub(
        r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")',
        r'\1Mercury Text\2', blob)

    theme_part._blob = blob.encode("utf-8")  # noqa: SLF001


def _label_master_footer(prs: Presentation, text: str = "Forge Brand · Confidential") -> None:
    """Set footer text on the first slide master's footer placeholder.

    The default Office master already carries ftr/dt/sldNum placeholders (which
    is what validate_deck's P0-4 identity check looks for); we just put brand
    text into the footer so it reads as intentional identity rather than empty.
    """
    master = prs.slide_masters[0]
    for ph in master.placeholders:
        if ph.placeholder_format.type is not None and "FOOTER" in str(ph.placeholder_format.type):
            if hasattr(ph, "text"):
                ph.text = text
            return
    # No footer placeholder found; that's fine — the layouts still carry one.


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("out", type=Path, help="output .pptx path")
    args = p.parse_args()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _rewrite_theme_xml(prs)
    _label_master_footer(prs)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    print(f"wrote {args.out} ({len(prs.slide_layouts)} layouts)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
